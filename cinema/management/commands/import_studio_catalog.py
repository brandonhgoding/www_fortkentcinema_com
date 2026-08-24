"""Import a studio's DCP availability list into StudioTitle.

Studios send a spreadsheet of everything they will rent us. This loads one
into the catalog behind the /sponsor/ page.

Idempotent: rows are matched on (studio, original_title, release_date), so
re-running after a studio ships an updated list adds the new titles and
refreshes the old ones without duplicating anything. `is_available` and
`hidden_reason` are deliberately NOT refreshed — a title you hid must stay
hidden across the next catalog drop.

Safe to run on prod:
    python manage.py import_studio_catalog --studio universal \\
        --file "Universal DCP List - Updated 3.4.26.xlsx"

Use --dry-run first to see the counts without writing.
"""

from pathlib import Path
from typing import Any

from cinema.models import Studio, StudioTitle
from cinema.studio_catalog import clean_cell, normalize_title, parse_runtime
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

# Slug -> display name. A studio must be listed here before its list can be
# imported, so a typo in --studio fails loudly instead of silently creating
# a "Univrsal" catalog nobody can find.
KNOWN_STUDIOS = {
    "universal": "Universal Pictures",
}

DEFAULT_SHEET = "by TITLE"

# Column order in every studio list we have seen so far.
COL_TITLE = 0
COL_RELEASE = 1
COL_ASPECT = 2
COL_AUDIO = 3
COL_RATING = 4
COL_RUNTIME = 5
COL_NOTES = 6


class ParsedRow:
    """One spreadsheet row, cleaned up."""

    def __init__(self, cells: tuple) -> None:
        def cell(index: int) -> Any:
            return cells[index] if index < len(cells) else None

        raw_title = clean_cell(cell(COL_TITLE))
        self.original_title = raw_title
        self.title = normalize_title(raw_title)

        release = cell(COL_RELEASE)
        self.release_date = release.date() if hasattr(release, "date") else None

        self.aspect_ratio = clean_cell(cell(COL_ASPECT))
        self.audio = clean_cell(cell(COL_AUDIO))
        self.mpa_rating = clean_cell(cell(COL_RATING))
        self.runtime_minutes, self.runtime_raw = parse_runtime(cell(COL_RUNTIME))
        self.notes = clean_cell(cell(COL_NOTES))

    @property
    def key(self) -> tuple[str, object]:
        return (self.original_title, self.release_date)

    @property
    def filled(self) -> int:
        """How many optional fields carry a value — the variant tie-break."""
        return sum(
            1
            for value in (
                self.aspect_ratio, self.audio, self.mpa_rating,
                self.runtime_raw, self.notes,
            )
            if value
        )

    def summary(self) -> str:
        """One-line description of this variant, for `variant_note`."""
        bits = [
            self.notes,
            f"{self.runtime_minutes} min" if self.runtime_minutes else "",
            self.audio,
            self.mpa_rating,
        ]
        return ", ".join(bit for bit in bits if bit)


class Command(BaseCommand):
    help = "Import a studio's DCP availability spreadsheet into StudioTitle."

    def add_arguments(self, parser) -> None:
        parser.add_argument(
            "--studio", required=True,
            help=f"Studio slug. One of: {', '.join(sorted(KNOWN_STUDIOS))}",
        )
        parser.add_argument("--file", required=True, help="Path to the .xlsx file.")
        parser.add_argument(
            "--sheet", default=DEFAULT_SHEET,
            help=(
                "Worksheet to read. Universal ships 'by DATE' and 'by TITLE' "
                "with identical content, so the default is fine."
            ),
        )
        parser.add_argument(
            "--dry-run", action="store_true",
            help="Report what would change without writing.",
        )

    def handle(self, *args: Any, **options: Any) -> None:
        slug = options["studio"]
        if slug not in KNOWN_STUDIOS:
            raise CommandError(
                f"Unknown studio '{slug}'. "
                f"Known: {', '.join(sorted(KNOWN_STUDIOS))}"
            )

        path = Path(options["file"])
        if not path.is_file():
            raise CommandError(f"No such file: {path}")

        rows = self._read(path, options["sheet"])
        groups = self._group(rows)

        dry_run = options["dry_run"]
        if dry_run:
            self.stdout.write(self.style.WARNING("DRY RUN — nothing written."))
            self.stdout.write(
                f"Would import {len(groups)} titles from {len(rows)} rows "
                f"({len(rows) - len(groups)} variants collapsed)."
            )
            return

        with transaction.atomic():
            studio, _ = Studio.objects.get_or_create(
                slug=slug, defaults={"name": KNOWN_STUDIOS[slug]}
            )
            created, updated, collapsed = self._write(studio, groups)

        self.stdout.write(
            self.style.SUCCESS(
                f"{studio.name}: {created} created, {updated} updated, "
                f"{collapsed} variants collapsed, {len(rows)} rows read."
            )
        )

    def _read(self, path: Path, sheet: str) -> list[ParsedRow]:
        # Imported here so the module loads even if openpyxl is missing,
        # which keeps `manage.py help` working with a clearer error.
        import openpyxl

        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if sheet not in book.sheetnames:
            raise CommandError(
                f"No sheet named '{sheet}' in {path.name}. "
                f"Found: {', '.join(book.sheetnames)}"
            )
        worksheet = book[sheet]
        rows = []
        for cells in worksheet.iter_rows(min_row=2, values_only=True):
            if not cells or not clean_cell(cells[COL_TITLE]):
                continue
            rows.append(ParsedRow(cells))
        book.close()
        return rows

    def _group(self, rows: list[ParsedRow]) -> dict:
        """Bucket rows by natural key, preserving sheet order within a bucket."""
        groups: dict = {}
        for row in rows:
            groups.setdefault(row.key, []).append(row)
        return groups

    def _primary(self, group: list[ParsedRow]) -> ParsedRow:
        """Pick the row that best represents a group of DCP variants.

        Longest runtime wins (the roadshow cut, the reconstruction). Ties go
        to the most complete row, then to sheet order — without that last
        step the two Dracula 1931-02-13 rows, identical but for their audio,
        would import differently run to run.
        """
        best = group[0]
        for row in group[1:]:
            if (row.runtime_minutes or 0) > (best.runtime_minutes or 0):
                best = row
            elif (row.runtime_minutes or 0) == (best.runtime_minutes or 0):
                if row.filled > best.filled:
                    best = row
        return best

    def _write(self, studio: Studio, groups: dict) -> tuple[int, int, int]:
        created = updated = collapsed = 0

        for group in groups.values():
            primary = self._primary(group)
            others = [row for row in group if row is not primary]
            collapsed += len(others)

            # Skip any variant that describes the same thing as the primary.
            # Notting Hill is listed twice, identically — without this the
            # page would advertise an "also available" version that is just
            # the row itself.
            seen = {primary.summary()}
            summaries = []
            for row in others:
                summary = row.summary()
                if not summary or summary in seen:
                    continue
                seen.add(summary)
                summaries.append(summary)

            variant_note = (
                f"Also available: {'; '.join(summaries)}"[:300] if summaries else ""
            )

            _, was_created = StudioTitle.objects.update_or_create(
                studio=studio,
                original_title=primary.original_title,
                release_date=primary.release_date,
                defaults={
                    "title": primary.title,
                    "aspect_ratio": primary.aspect_ratio,
                    "audio": primary.audio,
                    "mpa_rating": primary.mpa_rating,
                    "runtime_minutes": primary.runtime_minutes,
                    "runtime_raw": primary.runtime_raw,
                    "notes": primary.notes,
                    "variant_note": variant_note,
                    # NOTE: is_available / hidden_reason are absent on purpose.
                    # A re-import must never un-hide a title.
                },
            )
            if was_created:
                created += 1
            else:
                updated += 1

        return created, updated, collapsed
